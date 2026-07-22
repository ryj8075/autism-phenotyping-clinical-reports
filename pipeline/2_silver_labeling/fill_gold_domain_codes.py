#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
from collections import Counter
from pathlib import Path
from typing import Any

import openpyxl
import yaml

DEFAULT_INPUT = Path(__file__).parent / "data" / "gold_label" / "gold_labeling_template.xlsx"
DEFAULT_OUTPUT = DEFAULT_INPUT
DEFAULT_RULES = Path(__file__).parent / "config_local" / "gold_domain_rules.yaml"
DEFAULT_SHEET = "gold_labeling"
FALLBACK_DOMAIN = "C2"
DEFAULT_MAX_LABELS = 3

def _load_rule_file(path: Path) -> dict[str, Any]:
    if not path.exists():
        raise FileNotFoundError(
            f"Rule file not found: {path}. Provide --rules-file with a private YAML or JSON file."
        )

    with open(path, encoding="utf-8") as f:
        if path.suffix.lower() == ".json":
            data = json.load(f)
        else:
            data = yaml.safe_load(f)

    if not isinstance(data, dict):
        raise ValueError("Rule file must contain a mapping.")
    if "domain_rules" not in data or not isinstance(data["domain_rules"], dict):
        raise ValueError("Rule file must define a 'domain_rules' mapping.")
    return data

def _pairs(items: Any) -> list[tuple[str, float]]:
    out: list[tuple[str, float]] = []
    for item in items or []:
        if isinstance(item, dict):
            pattern = item.get("pattern")
            weight = item.get("weight", 1)
        elif isinstance(item, (list, tuple)) and len(item) == 2:
            pattern, weight = item
        else:
            continue
        if pattern:
            out.append((str(pattern), float(weight)))
    return out

def compile_rules(raw_rules: dict[str, Any]) -> dict[str, dict[str, list[tuple[re.Pattern, float]]]]:
    compiled: dict[str, dict[str, list[tuple[re.Pattern, float]]]] = {}
    for domain, spec in raw_rules.items():
        if not isinstance(spec, dict):
            continue
        compiled[str(domain)] = {
            "keywords": [(re.compile(p), w) for p, w in _pairs(spec.get("keywords"))],
            "negative": [(re.compile(p), w) for p, w in _pairs(spec.get("negative"))],
        }
    return compiled

def score_sentence(
    sentence: str,
    rules: dict[str, dict[str, list[tuple[re.Pattern, float]]]],
) -> dict[str, float]:
    scores: dict[str, float] = {domain: 0.0 for domain in rules}
    for domain, spec in rules.items():
        for pattern, weight in spec.get("keywords", []):
            if pattern.search(sentence):
                scores[domain] += weight
        for pattern, weight in spec.get("negative", []):
            if pattern.search(sentence):
                scores[domain] += weight
    return scores

def select_domains(
    scores: dict[str, float],
    priority: list[str],
    fallback: str = FALLBACK_DOMAIN,
    max_labels: int = DEFAULT_MAX_LABELS,
    min_score: float = 1.0,
) -> list[str]:
    rank = {domain: i for i, domain in enumerate(priority)}
    candidates = [(domain, score) for domain, score in scores.items() if score >= min_score]
    candidates.sort(key=lambda item: (-item[1], rank.get(item[0], 999)))
    selected = [domain for domain, _ in candidates[:max_labels]]
    return selected or [fallback]

def classify_sentence(
    sentence: str,
    rules: dict[str, dict[str, list[tuple[re.Pattern, float]]]],
    priority: list[str],
    fallback: str,
    max_labels: int,
    min_score: float,
) -> str:
    if not sentence or not sentence.strip():
        return fallback
    scores = score_sentence(sentence, rules)
    return ",".join(select_domains(scores, priority, fallback, max_labels, min_score))

def fill_workbook(
    input_path: Path,
    output_path: Path,
    rules_path: Path,
    sheet_name: str,
    fallback: str,
    max_labels: int,
    min_score: float,
) -> None:
    rule_data = _load_rule_file(rules_path)
    rules = compile_rules(rule_data["domain_rules"])
    priority = [str(x) for x in rule_data.get("domain_priority", list(rules.keys()))]

    wb = openpyxl.load_workbook(str(input_path))
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Sheet not found: {sheet_name}. Available sheets: {wb.sheetnames}")
    ws = wb[sheet_name]

    sentences_data = []
    for row in range(2, ws.max_row + 1):
        report_id = ws.cell(row=row, column=1).value
        sent_idx = ws.cell(row=row, column=2).value
        sentence = ws.cell(row=row, column=3).value
        sentences_data.append((row, report_id, sent_idx, sentence))

    domain_counts: Counter[str] = Counter()
    label_card_counts: Counter[int] = Counter()

    for row, _, _, sentence in sentences_data:
        domain_code = classify_sentence(
            str(sentence or ""),
            rules,
            priority,
            fallback,
            max_labels,
            min_score,
        )
        ws.cell(row=row, column=4, value=domain_code)
        codes = [code for code in domain_code.split(",") if code]
        domain_counts.update(codes)
        label_card_counts[len(codes)] += 1

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output_path))

    n_sent = len(sentences_data)
    print(f"Loaded sentences: {n_sent}")
    print(f"Saved workbook: {output_path}")
    print("Domain distribution:")
    for code in sorted(domain_counts):
        count = domain_counts[code]
        pct = count / n_sent * 100 if n_sent else 0
        print(f"  {code}: {count:4d} ({pct:5.1f}% of sentences)")
    print("Labels per sentence:")
    for k in sorted(label_card_counts):
        print(f"  {k} labels: {label_card_counts[k]:4d} sentences")

def main() -> None:
    parser = argparse.ArgumentParser(
        description="Fill a gold-label workbook's domain_code column from a private rule file."
    )
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--rules-file", type=Path, default=DEFAULT_RULES)
    parser.add_argument("--sheet-name", default=DEFAULT_SHEET)
    parser.add_argument("--fallback", default=FALLBACK_DOMAIN)
    parser.add_argument("--max-labels", type=int, default=DEFAULT_MAX_LABELS)
    parser.add_argument("--min-score", type=float, default=1.0)
    args = parser.parse_args()

    fill_workbook(
        input_path=args.input,
        output_path=args.output,
        rules_path=args.rules_file,
        sheet_name=args.sheet_name,
        fallback=args.fallback,
        max_labels=args.max_labels,
        min_score=args.min_score,
    )

if __name__ == "__main__":
    main()
