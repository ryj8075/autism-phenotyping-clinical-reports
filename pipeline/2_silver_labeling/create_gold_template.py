import argparse
import json
import os
import sys
from pathlib import Path
from typing import Any, Dict, List, Optional

import openpyxl
import yaml
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
REPORTS_DIR = BASE_DIR.parent / "data" / "all1st309" / "reports_txt_no_headings"
OUTPUT_PATH = DATA_DIR / "outputs" / "gold_labeling_template.xlsx"
GOLD_JSON = DATA_DIR / "gold_target_reports.json"

DEFAULT_DOMAINS_FILE = BASE_DIR / "domains_19.yaml"

# segment.py import
sys.path.insert(0, str(BASE_DIR))
from segment import segment_sentences

def load_domains_for_template(
    config_path: Optional[str] = None,
    domains_file: Optional[str] = None,
    pilot_mode: bool = False,
) -> List[Dict[str, Any]]:

    if not domains_file:
        if config_path:
            with open(config_path, encoding="utf-8") as f:
                cfg = yaml.safe_load(f)
            domains_file = cfg.get("domains_file")
            if domains_file and not Path(domains_file).is_absolute():
                domains_file = str(Path(config_path).parent / domains_file)

    if not domains_file:
        raise FileNotFoundError(
            "Domain file was not found.\n"
            "  Set --domains-file or config.yaml domains_file."
        )

    with open(domains_file, encoding="utf-8") as f:
        dcfg = yaml.safe_load(f)

    if not pilot_mode:
        pilot_mode = dcfg.get("pilot_mode", False)

    all_domains = []
    for category in dcfg.get("categories", []):
        cat_name = category.get("name", "")
        for subcat in category.get("subcategories", []):
            for domain in subcat.get("domains", []):
                all_domains.append({
                    "code": str(domain["code"]),
                    "id": domain["id"],
                    "name_ko": domain["name_ko"],
                    "name_en": domain["name_en"],
                    "description": domain.get("description", "").strip(),
                    "pilot_active": domain.get("pilot_active", True),
                    "category": cat_name,
                })

    if pilot_mode:
        active = [d for d in all_domains if d["pilot_active"]]
    else:
        active = all_domains

    print(f"Loaded domains: total {len(all_domains)}, active {len(active)} (pilot_mode={pilot_mode})")
    return active

def main():
    parser = argparse.ArgumentParser(
        description="Create an expert gold-labeling xlsx template (multi-label, 19 domains)",
    )
    parser.add_argument(
        "--config", "-c",
        default=str(BASE_DIR / "config.yaml"),
        help="Configuration file path (default: config.yaml)",
    )
    parser.add_argument(
        "--domains-file",
        default=str(DEFAULT_DOMAINS_FILE),
        help=f"Domain-definition YAML path (default: {DEFAULT_DOMAINS_FILE}, 19 domains)",
    )
    parser.add_argument(
        "--pilot-mode",
        action="store_true",
        default=False,
        help="Pilot mode: use A+C domains only (default: all 19 domains)",
    )
    parser.add_argument(
        "--reports-dir",
        default=str(REPORTS_DIR),
        help=f"Directory containing report txt files (default: {REPORTS_DIR})",
    )
    parser.add_argument(
        "--gold-json",
        default=str(GOLD_JSON),
        help=f"JSON list of reports to include (default: {GOLD_JSON})",
    )
    parser.add_argument(
        "--output", "-o",
        default=str(OUTPUT_PATH),
        help=f"Output xlsx path (default: {OUTPUT_PATH})",
    )
    args = parser.parse_args()

    reports_dir = Path(args.reports_dir)
    gold_json = Path(args.gold_json)

    domains = load_domains_for_template(
        config_path=args.config,
        domains_file=args.domains_file,
        pilot_mode=args.pilot_mode,
    )

    with open(args.config, encoding="utf-8") as f:
        cfg = yaml.safe_load(f)
    seg_cfg = cfg.get("segmentation", {})
    seg_method = seg_cfg.get("method", "tokenizer_aligned")
    seg_min = seg_cfg.get("min_length", 30)
    seg_max = seg_cfg.get("max_length", 500)

    with open(gold_json, encoding="utf-8") as f:
        gold_data = json.load(f)
    report_files = gold_data["reports"]
    print(f"Loaded {len(report_files)} target reports")

    wb = openpyxl.Workbook()

    header_font = Font(bold=True, size=12)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font_white = Font(bold=True, size=12, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    title_font = Font(bold=True, size=14)

    # ────────────────────────────────────────────────────────────

    # ────────────────────────────────────────────────────────────
    ws_guide = wb.active
    ws_guide.title = "Labeling Guide"

    ws_guide.merge_cells("A1:E1")
    mode_label = "Pilot Mode" if args.pilot_mode else "All Domains"
    ws_guide["A1"] = f"ASD Clinical Report Domain Labeling Guide ({mode_label}, {len(domains)} domains)"
    ws_guide["A1"].font = title_font
    ws_guide["A1"].alignment = Alignment(horizontal="center")

    row = 3

    guide_headers = ["Code", "Category", "Domain Name (legacy field)", "Domain Name (English)", "Description"]
    for col_idx, header_text in enumerate(guide_headers, start=1):
        cell = ws_guide.cell(row=row, column=col_idx, value=header_text)
        cell.font = header_font_white
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for i, d in enumerate(domains):
        r = row + 1 + i
        ws_guide.cell(row=r, column=1, value=d["code"]).border = thin_border
        ws_guide.cell(row=r, column=1).alignment = Alignment(horizontal="center")
        ws_guide.cell(row=r, column=2, value=d["category"]).border = thin_border
        ws_guide.cell(row=r, column=3, value=d["name_ko"]).border = thin_border
        ws_guide.cell(row=r, column=4, value=d["name_en"]).border = thin_border
        ws_guide.cell(row=r, column=5, value=d["description"]).border = thin_border
        ws_guide.cell(row=r, column=5).alignment = Alignment(wrap_text=True)

        if i % 2 == 1:
            light_fill = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
            for c in range(1, 6):
                ws_guide.cell(row=r, column=c).fill = light_fill

    rule_start = row + 1 + len(domains) + 1
    ws_guide.merge_cells(f"A{rule_start}:E{rule_start}")
    ws_guide.cell(row=rule_start, column=1, value="Labeling Rules").font = Font(bold=True, size=13)

    fallback_code = "C2"
    for d in domains:
        if d["id"] == "other_general":
            fallback_code = d["code"]
            break

    rules = [
        "1. Enter the applicable domain code in column D (domain_code) of the gold_labeling sheet.",
        "2. If multiple domains apply, enter all of them separated by commas (for example: A1, B3).",
        "3. Put the primary domain first; order reflects priority.",
        f"4. If no domain applies, enter {fallback_code} (Other / General).",
        "5. Use the Code column in the table above.",
        "6. Blank cells are treated as missing labels.",
    ]

    for i, rule_text in enumerate(rules):
        r = rule_start + 1 + i
        ws_guide.merge_cells(f"A{r}:E{r}")
        ws_guide.cell(row=r, column=1, value=rule_text).font = Font(size=11)

    ws_guide.column_dimensions["A"].width = 10
    ws_guide.column_dimensions["B"].width = 30
    ws_guide.column_dimensions["C"].width = 28
    ws_guide.column_dimensions["D"].width = 32
    ws_guide.column_dimensions["E"].width = 60

    # ────────────────────────────────────────────────────────────

    # ────────────────────────────────────────────────────────────
    ws_label = wb.create_sheet(title="gold_labeling")

    label_headers = ["report_id", "sentence_idx", "sentence", "domain_code"]
    for col_idx, header_text in enumerate(label_headers, start=1):
        cell = ws_label.cell(row=1, column=col_idx, value=header_text)
        cell.font = header_font_white
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center")

    fill_light_blue = PatternFill(start_color="DAEEF3", end_color="DAEEF3", fill_type="solid")
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    current_row = 2
    total_sentences = 0

    for report_idx, report_file in enumerate(report_files):
        report_id = os.path.splitext(report_file)[0]
        report_path = reports_dir / report_file

        if not report_path.exists():
            print(f"  [warning] File not found: {report_path}")
            continue

        with open(report_path, encoding="utf-8") as f:
            text = f.read()

        sentences = segment_sentences(
            text,
            method=seg_method,
            min_length=seg_min,
            max_length=seg_max,
        )
        print(f"  [{report_idx+1:2d}/{len(report_files)}] {report_id}: {len(sentences)} sentences")

        row_fill = fill_light_blue if report_idx % 2 == 0 else fill_white

        for sent_idx, sentence in enumerate(sentences):
            r = current_row

            cell_a = ws_label.cell(row=r, column=1, value=report_id)
            cell_b = ws_label.cell(row=r, column=2, value=sent_idx)
            cell_c = ws_label.cell(row=r, column=3, value=sentence)
            cell_d = ws_label.cell(row=r, column=4, value="")

            for cell in [cell_a, cell_b, cell_c, cell_d]:
                cell.fill = row_fill
                cell.border = thin_border

            cell_a.alignment = Alignment(vertical="center")
            cell_b.alignment = Alignment(horizontal="center", vertical="center")
            cell_c.alignment = Alignment(wrap_text=True, vertical="top")
            cell_d.alignment = Alignment(horizontal="center", vertical="center")

            current_row += 1
            total_sentences += 1

    ws_label.column_dimensions["A"].width = 22
    ws_label.column_dimensions["B"].width = 14
    ws_label.column_dimensions["C"].width = 80
    ws_label.column_dimensions["D"].width = 22

    ws_label.freeze_panes = "A2"

    output_path = Path(args.output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output_path))
    print(f"\nCreated file: {output_path}")
    print(f"Total reports: {len(report_files)}, total sentences: {total_sentences}")
    print(f"Domains: {len(domains)} ({'pilot' if args.pilot_mode else 'all'})")
    print(f"Segmentation: method={seg_method}, min_length={seg_min}")
    print(f"File size: {os.path.getsize(str(output_path)):,} bytes")

if __name__ == "__main__":
    main()
