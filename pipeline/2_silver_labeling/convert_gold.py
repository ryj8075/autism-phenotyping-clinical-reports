from __future__ import annotations

import argparse
import json
import os
import sys
from collections import Counter
from pathlib import Path
from typing import Any, Dict, List, Optional

import openpyxl
import yaml

BASE_DIR = Path(__file__).resolve().parent
DEFAULT_INPUT = BASE_DIR / "data" / "outputs" / "gold_labeling_template.xlsx"
DEFAULT_OUTPUT = BASE_DIR / "data" / "outputs" / "gold_labels.jsonl"
SHEET_NAME = "gold_labeling"

def load_domains_for_convert(
    config_path: Optional[str] = None,
    domains_file: Optional[str] = None,
    pilot_mode: bool = False,
) -> List[Dict[str, Any]]:

    if not domains_file:
        if config_path:
            with open(config_path, encoding="utf-8") as f:
                cfg = yaml.safe_load(f)
            domains_file = cfg.get("domains_file")
            if domains_file and not Path(domains_file).is_absolute():
                domains_file = str(Path(config_path).parent / domains_file)

    if not domains_file:
        raise FileNotFoundError(
            "Domain file was not found.\n"
            "  Set --domains-file or config.yaml domains_file."
        )

    with open(domains_file, encoding="utf-8") as f:
        dcfg = yaml.safe_load(f)

    if not pilot_mode:
        pilot_mode = dcfg.get("pilot_mode", False)

    all_domains = []
    for category in dcfg.get("categories", []):
        cat_name = category.get("name", "")
        for subcat in category.get("subcategories", []):
            for domain in subcat.get("domains", []):
                all_domains.append({
                    "code": str(domain["code"]),
                    "id": domain["id"],
                    "name_ko": domain["name_ko"],
                    "name_en": domain["name_en"],
                    "description": domain.get("description", "").strip(),
                    "pilot_active": domain.get("pilot_active", True),
                    "category": cat_name,
                })

    if pilot_mode:
        active = [d for d in all_domains if d["pilot_active"]]
    else:
        active = all_domains

    print(f"Loaded domains: total {len(all_domains)}, active {len(active)} (pilot_mode={pilot_mode})")
    return active

def build_code_lookup(domains: List[Dict[str, Any]]) -> Dict[str, Dict[str, Any]]:
    lookup = {}
    for d in domains:
        lookup[d["code"].upper()] = d
        lookup[d["code"].lower()] = d
        lookup[d["id"]] = d
    return lookup

def _lookup_token(token: str, code_lookup: Dict[str, Dict[str, Any]]) -> Optional[str]:
    if token in code_lookup:
        return code_lookup[token]["id"]

    if token.upper() in code_lookup:
        return code_lookup[token.upper()]["id"]

    return None

def parse_domain_code(
    raw_value,
    code_lookup: Dict[str, Dict[str, Any]],
    warnings: Optional[list] = None,
    row_num: Optional[int] = None,
    report_id: str = "",
) -> List[str]:

    if raw_value is None:
        return []

    if isinstance(raw_value, (int, float)):
        text = str(int(raw_value)).strip()
    else:
        text = str(raw_value).strip()

    if not text:
        return []

    domain_ids: List[str] = []
    for token in text.split(","):
        token = token.strip()
        if not token:
            continue
        did = _lookup_token(token, code_lookup)
        if did is None:

            if warnings is not None:
                msg = f"[warning] Row {row_num} ({report_id}): unrecognized domain code '{token}'"
                warnings.append(msg)
            continue
        if did not in domain_ids:
            domain_ids.append(did)

    return domain_ids

def validate_code(
    domain_id: Optional[str],
    row_num: int,
    report_id: str,
    valid_ids: set,
    warnings: list,
) -> Optional[str]:

    if domain_id is None:
        return None
    if domain_id in valid_ids:
        return domain_id
    msg = f"[warning] Row {row_num} ({report_id}): ID not in active domains '{domain_id}'"
    warnings.append(msg)
    return None

def convert(
    input_path: str,
    output_path: str,
    domains: List[Dict[str, Any]],
) -> None:

    if not os.path.exists(input_path):
        print(f"[error] Input file not found: {input_path}")
        sys.exit(1)

    print(f"Input file: {input_path}")
    print(f"Output file: {output_path}")
    print(f"Number of domains: {len(domains)}")
    print()

    code_lookup = build_code_lookup(domains)
    valid_ids = {d["id"] for d in domains}

    fallback_id = "other_general"
    fallback_code = "RE2"
    for d in domains:
        if d["id"] == "other_general":
            fallback_id = d["id"]
            fallback_code = d["code"]
            break

    wb = openpyxl.load_workbook(input_path, read_only=True, data_only=True)

    if SHEET_NAME not in wb.sheetnames:
        print(f"[error] Sheet '{SHEET_NAME}' was not found.")
        print(f"  Available sheets: {wb.sheetnames}")
        sys.exit(1)

    ws = wb[SHEET_NAME]
    print(f"Loaded sheet '{SHEET_NAME}'")

    rows = ws.iter_rows(min_row=1, max_row=1, values_only=True)
    header = next(rows)
    expected_header = ("report_id", "sentence_idx", "sentence", "domain_code")

    for i, (got, want) in enumerate(zip(header, expected_header)):
        if got != want:
            print(f"[warning] Column {i+1} header mismatch: '{got}' (expected: '{want}')")

    records = []
    warnings = []
    empty_label_count = 0
    multi_label_count = 0
    total_rows = 0
    domain_counter = Counter()

    for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row[0] is None and row[2] is None:
            continue

        report_id = str(row[0]) if row[0] is not None else ""
        sentence_idx = int(row[1]) if row[1] is not None else 0
        sentence = str(row[2]) if row[2] is not None else ""
        raw_code = row[3]

        total_rows += 1

        parsed_ids = parse_domain_code(
            raw_code, code_lookup,
            warnings=warnings, row_num=row_num, report_id=report_id,
        )

        labels: List[str] = []
        for did in parsed_ids:
            valid = validate_code(did, row_num, report_id, valid_ids, warnings)
            if valid is not None and valid not in labels:
                labels.append(valid)

        if not labels:

            empty_label_count += 1
            labels = [fallback_id]
        elif len(labels) >= 2:
            multi_label_count += 1

        for did in labels:
            domain_counter[did] += 1

        record = {
            "report_id": report_id,
            "sentence_idx": sentence_idx,
            "sentence": sentence,
            "labels": labels,
        }
        records.append(record)

    wb.close()

    os.makedirs(os.path.dirname(output_path), exist_ok=True)

    with open(output_path, "w", encoding="utf-8") as f:
        for rec in records:
            f.write(json.dumps(rec, ensure_ascii=False) + "\n")

    if warnings:
        print(f"\n{'='*72}")
        print(f"Warnings ({len(warnings)}):")
        print(f"{'='*72}")
        for w in warnings:
            print(f"  {w}")

    labeled_count = total_rows - empty_label_count
    total_label_assignments = sum(domain_counter.values())
    avg_labels = total_label_assignments / total_rows if total_rows else 0.0

    print(f"\n{'='*72}")
    print(f"Conversion summary (Multi-label, {len(domains)} domains)")
    print(f"{'='*72}")
    print(f"  Total sentences:          {total_rows:,}")
    print(f"  Labeled sentences:        {labeled_count:,} ({labeled_count/total_rows*100:.1f}%)" if total_rows else "  Labeled sentences:        0")
    print(f"  Empty-label sentences:    {empty_label_count:,} ({empty_label_count/total_rows*100:.1f}%)" if total_rows else "  Empty-label sentences:    0")
    print(f"  Multi-label sentences:    {multi_label_count:,} ({multi_label_count/total_rows*100:.1f}%)" if total_rows else "  Multi-label sentences:    0")
    print(f"  Mean labels per sentence: {avg_labels:.2f}")
    print(f"  Empty-label fallback:     {fallback_code} ({fallback_id})")

    print(f"\n  Domain distribution (all labels per sentence):")
    for d in domains:
        code = d["code"]
        did = d["id"]
        name = d["name_ko"]
        count = domain_counter.get(did, 0)
        bar = "#" * min(count, 50)
        pct = f"({count/total_rows*100:.1f}%)" if total_rows else ""
        print(f"    {code:>3}  {name:<20s}  {count:>5,}  {pct:>8s}  {bar}")

    print(f"\n  Output file: {output_path}")
    print(f"  File size: {os.path.getsize(output_path):,} bytes")
    print(f"{'='*72}")

# ── CLI ───────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(
        description="Convert expert gold-labeling xlsx files to gold_labels.jsonl (multi-label, 19 domains)",
    )
    parser.add_argument(
        "--config", "-c",
        default=str(BASE_DIR / "config.yaml"),
        help="Configuration file path (default: config.yaml)",
    )
    parser.add_argument(
        "--domains-file",
        default=None,
        help="Domain-definition YAML path (overrides config domains_file)",
    )
    parser.add_argument(
        "--pilot-mode",
        action="store_true",
        default=False,
        help="Pilot mode: use A+C domains only (default: all 19 domains)",
    )
    parser.add_argument(
        "--input", "-i",
        default=str(DEFAULT_INPUT),
        help=f"Input xlsx path (default: {DEFAULT_INPUT})",
    )
    parser.add_argument(
        "--output", "-o",
        default=str(DEFAULT_OUTPUT),
        help=f"Output jsonl path (default: {DEFAULT_OUTPUT})",
    )
    args = parser.parse_args()

    domains = load_domains_for_convert(
        config_path=args.config,
        domains_file=args.domains_file,
        pilot_mode=args.pilot_mode,
    )

    convert(args.input, args.output, domains)

if __name__ == "__main__":
    main()
